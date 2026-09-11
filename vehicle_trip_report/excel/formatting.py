from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Fills (ARGB format)
HEADER_FILL_PRATAPGARH = PatternFill(start_color="FF366092", end_color="FF366092", fill_type="solid")
HEADER_FILL_MUJERI = PatternFill(start_color="FF366092", end_color="FF366092", fill_type="solid")
HEADER_FILL_GURGAON = PatternFill(start_color="FF366092", end_color="FF366092", fill_type="solid")
SUMMARY_HEADER_FILL = PatternFill(start_color="FF366092", end_color="FF366092", fill_type="solid")

# Fonts
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="FF000000")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFFFF")
DATA_FONT = Font(name="Arial", size=11)
BOLD_FONT = Font(name="Arial", size=11, bold=True)

# Alignments
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)

# Borders
THIN_BORDER = Border(
    left=Side(style='thin', color="FFDDDDDD"),
    right=Side(style='thin', color="FFDDDDDD"),
    top=Side(style='thin', color="FFDDDDDD"),
    bottom=Side(style='thin', color="FFDDDDDD")
)

def set_column_widths(ws, columns, widths):
    for i, col_name in enumerate(columns, start=1):
        # Determine column letter
        from openpyxl.utils import get_column_letter
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = widths[i-1]

def apply_title_formatting(ws, end_col):
    from openpyxl.utils import get_column_letter
    end_col_letter = get_column_letter(end_col)
    
    # Merge first row
    ws.merge_cells(f'A1:{end_col_letter}1')
    title_cell = ws['A1']
    title_cell.font = TITLE_FONT
    title_cell.alignment = CENTER_ALIGN

def apply_header_formatting(ws, row_idx, fill_color):
    import copy
    for cell in ws[row_idx]:
        cell.font = copy.copy(HEADER_FONT)
        cell.fill = copy.copy(fill_color)
        cell.alignment = copy.copy(CENTER_ALIGN)
        cell.border = copy.copy(THIN_BORDER)

def apply_data_row_formatting(ws, row_idx):
    import copy
    for cell in ws[row_idx]:
        cell.font = copy.copy(DATA_FONT)
        cell.alignment = copy.copy(CENTER_ALIGN)
        cell.border = copy.copy(THIN_BORDER)
