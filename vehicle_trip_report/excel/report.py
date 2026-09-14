import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
import sys
from datetime import datetime

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config import REPORTS_DIR, ROW_HEIGHT, EXCEL_IMAGE_WIDTH, EXCEL_IMAGE_HEIGHT, UNIT_AGENCIES
from excel.formatting import (
    apply_title_formatting, apply_header_formatting, apply_data_row_formatting, 
    set_column_widths, HEADER_FILL_PRATAPGARH, HEADER_FILL_MUJERI, SUMMARY_HEADER_FILL,
    BOLD_FONT, CENTER_ALIGN, LEFT_ALIGN, THIN_BORDER
)
from utils.image_utils import process_image_for_excel

def get_report_path(report_name):
    if not report_name.endswith('.xlsx'):
        report_name += '.xlsx'
    return REPORTS_DIR / report_name

def generate_daily_report(date_str, records, slip_type="Input"):
    """
    date_str: DD-MM-YYYY
    records: list of dicts from db
    """
    dt = datetime.strptime(date_str, "%d-%m-%Y")
    formatted_date_title = dt.strftime("%d/%m/%Y")
    tab_date = dt.strftime("%d-%b")
    
    report_name = f"Vehicle_Trip_Report_{date_str}.xlsx"
    report_path = get_report_path(report_name)
    
    wb = Workbook()
    
    default_sheet = wb.active
    
    df = pd.DataFrame(records)
    if df.empty:
        return None
        
    # Sort by RST No in ascending order
    if 'rst_no' in df.columns:
        df['rst_no_num'] = pd.to_numeric(df['rst_no'], errors='coerce')
        df = df.sort_values(by=['rst_no_num', 'rst_no'], ascending=[True, True])
        
    if 'processing_unit' not in df.columns:
        df['processing_unit'] = 'Pratapgarh'
    df['processing_unit'] = df['processing_unit'].fillna('Pratapgarh')
    
    # Fix missing or 'Other' agency names from older records
    df.loc[(df['processing_unit'] == 'Mujeri') & ((df['agency_name'].isnull()) | (df['agency_name'] == '') | (df['agency_name'] == 'Other')), 'agency_name'] = 'V S Waste'
    df.loc[(df['processing_unit'] == 'Pratapgarh') & ((df['agency_name'].isnull()) | (df['agency_name'] == '')), 'agency_name'] = 'SVN'
        
    units = df['processing_unit'].unique()
    
    for unit in ['Pratapgarh', 'Mujeri', 'Gurgaon Paper Mills']:
        if unit not in units:
            continue
            
        unit_df = df[df['processing_unit'] == unit].copy()
        
        # Ensure net_weight is numeric for summing and excel formatting
        unit_df['net_weight'] = pd.to_numeric(unit_df['net_weight'], errors='coerce').fillna(0)
        
        # 1. Data Sheet
        data_sheet = wb.create_sheet(f"{unit} {tab_date}")
        
        title = f"Municipal Corporation Faridabad - {unit} Processing Unit - {slip_type} Vehicle Trip Report ({formatted_date_title})"
        data_sheet.append([title])
        data_sheet.append([])
        
        headers = ["S.No", "RST No (Slip No.)", "Vehicle Number", "Net Weight (kg)", "Agency Name", "Reported By", "Slip Image"]
        widths = [8, 18, 20, 18, 20, 20, 45]
        
        if unit == "Pratapgarh":
            header_fill = HEADER_FILL_PRATAPGARH
        elif unit == "Mujeri":
            header_fill = HEADER_FILL_MUJERI
        else:
            from excel.formatting import HEADER_FILL_GURGAON
            header_fill = HEADER_FILL_GURGAON
            
        data_sheet.append(headers)
        
        apply_title_formatting(data_sheet, len(headers))
        apply_header_formatting(data_sheet, 3, header_fill)
        set_column_widths(data_sheet, headers, widths)
        
        next_row = 4
        s_no = 1
        for _, record in unit_df.iterrows():
            row_data = [
                s_no,
                record.get('rst_no', ''),
                record.get('vehicle_no', ''),
                record.get('net_weight', '')
            ]
            
            row_data.append(record.get('agency_name', ''))
            row_data.append(record.get('reported_by', ''))
            row_data.append("") # Placeholder for image
            
            data_sheet.append(row_data)
            
            # Embed Image
            image_path = record.get('image_path')
            if image_path and os.path.exists(image_path):
                temp_img_path = process_image_for_excel(image_path)
                if temp_img_path:
                    img = OpenpyxlImage(temp_img_path)
                    img_col_letter = 'G'
                    img.anchor = f"{img_col_letter}{next_row}"
                    data_sheet.add_image(img)
                    data_sheet.row_dimensions[next_row].height = ROW_HEIGHT
            else:
                img_col_idx = 7
                data_sheet.cell(row=next_row, column=img_col_idx, value="No Image")
                
            apply_data_row_formatting(data_sheet, next_row)
            next_row += 1
            s_no += 1
            
        # Add Total Row
        total_weight = unit_df['net_weight'].sum()
        total_row = [""] * len(headers)
        total_row[2] = "Total Net Weight (kg)"
        total_row[3] = total_weight
        
        data_sheet.append(total_row)
        for col_idx in range(1, len(headers) + 1):
            cell = data_sheet.cell(row=next_row, column=col_idx)
            cell.font = BOLD_FONT
            cell.alignment = CENTER_ALIGN
            cell.border = THIN_BORDER
            
        next_row += 1
            
        # 2. Summary Sheet
        summary_sheet = wb.create_sheet(f"{unit} Summary")
        
        summary_sheet.append([f"Summary - {unit} Processing Unit {slip_type} Vehicle Trip Report ({formatted_date_title})"])
        summary_sheet['A1'].font = BOLD_FONT
        summary_sheet.append([])
        
        current_row = 3
        
        # Table 1: Overall Statistics
        summary_sheet.cell(row=current_row, column=1, value="Overall Statistics").font = BOLD_FONT
        current_row += 1
        
        total_trips = len(unit_df)
        total_weight = unit_df['net_weight'].sum()
        avg_weight = total_weight / total_trips if total_trips > 0 else 0
        
        stats = [
            ["Total Trips", total_trips],
            ["Total Net Weight (kg)", total_weight],
            ["Average Net Weight per Trip", round(avg_weight, 5)]
        ]
        
        for row_data in stats:
            summary_sheet.append(row_data)
            apply_data_row_formatting(summary_sheet, current_row)
            summary_sheet.cell(row=current_row, column=1).alignment = LEFT_ALIGN
            current_row += 1
            
        current_row += 2
        
        # Table 2: Agency-wise Breakdown
        summary_sheet.cell(row=current_row, column=1, value="Agency-wise Breakdown").font = BOLD_FONT
        current_row += 1
        
        headers = ["Agency Name", "No. of Trips", "Total Net Weight (kg)", "% of Total Weight"]
        summary_sheet.append(headers)
        apply_header_formatting(summary_sheet, current_row, SUMMARY_HEADER_FILL)
        current_row += 1
        
        agency_grouped = unit_df.groupby('agency_name').agg(
            trips=('rst_no', 'count'),
            weight=('net_weight', 'sum')
        ).reset_index()
        
        for _, row in agency_grouped.iterrows():
            pct = row['weight'] / total_weight if total_weight > 0 else 0
            summary_sheet.append([row['agency_name'], row['trips'], row['weight'], pct])
            apply_data_row_formatting(summary_sheet, current_row)
            summary_sheet.cell(row=current_row, column=1).alignment = LEFT_ALIGN
            current_row += 1
            
        summary_sheet.append(["Total", total_trips, total_weight, 1])
        apply_data_row_formatting(summary_sheet, current_row)
        for col in range(1, 5):
            summary_sheet.cell(row=current_row, column=col).font = BOLD_FONT
        summary_sheet.cell(row=current_row, column=1).alignment = LEFT_ALIGN
        current_row += 2
            
        # Table 3: Reported By - Breakdown
        summary_sheet.cell(row=current_row, column=1, value="Reported By - Breakdown").font = BOLD_FONT
        current_row += 1
        
        headers = ["Reported By", "No. of Trips", "Total Net Weight (kg)"]
        summary_sheet.append(headers)
        apply_header_formatting(summary_sheet, current_row, SUMMARY_HEADER_FILL)
        current_row += 1
        
        reporter_grouped = unit_df.groupby('reported_by').agg(
            trips=('rst_no', 'count'),
            weight=('net_weight', 'sum')
        ).reset_index()
        
        for _, row in reporter_grouped.iterrows():
            summary_sheet.append([row['reported_by'], row['trips'], row['weight']])
            apply_data_row_formatting(summary_sheet, current_row)
            summary_sheet.cell(row=current_row, column=1).alignment = LEFT_ALIGN
            current_row += 1
            
        current_row += 2
        
        # Table 4: Vehicle-wise Trip Count
        summary_sheet.cell(row=current_row, column=1, value="Vehicle-wise Trip Count (Repeat Vehicles)").font = BOLD_FONT
        current_row += 1
        
        headers = ["Vehicle Number", "No. of Trips", "Total Net Weight (kg)"]
        summary_sheet.append(headers)
        apply_header_formatting(summary_sheet, current_row, SUMMARY_HEADER_FILL)
        current_row += 1
        
        vehicle_grouped = unit_df.groupby('vehicle_no').agg(
            trips=('rst_no', 'count'),
            weight=('net_weight', 'sum')
        ).reset_index().sort_values('trips', ascending=False)
        
        for _, row in vehicle_grouped.iterrows():
            summary_sheet.append([row['vehicle_no'], row['trips'], row['weight']])
            apply_data_row_formatting(summary_sheet, current_row)
            summary_sheet.cell(row=current_row, column=1).alignment = LEFT_ALIGN
            current_row += 1
            
        set_column_widths(summary_sheet, ["A", "B", "C", "D"], [35, 15, 25, 20])

    if len(wb.sheetnames) > 1 and default_sheet.title in wb.sheetnames:
        wb.remove(default_sheet)
        
    wb.save(report_path)
    return str(report_path)

def generate_monthly_summary(month_year_str, records, slip_type="Input"):
    """
    month_year_str: MM-YYYY
    records: list of dicts from db
    """
    report_name = f"Monthly_Summary_{month_year_str}.xlsx"
    report_path = get_report_path(report_name)
    
    wb = Workbook()
    default_sheet = wb.active
    
    df = pd.DataFrame(records)
    if df.empty:
        return None
        
    if 'processing_unit' not in df.columns:
        df['processing_unit'] = 'Pratapgarh'
    df['processing_unit'] = df['processing_unit'].fillna('Pratapgarh')
    
    # Fix missing or 'Other' agency names from older records
    df.loc[(df['processing_unit'] == 'Mujeri') & ((df['agency_name'].isnull()) | (df['agency_name'] == '') | (df['agency_name'] == 'Other')), 'agency_name'] = 'V S Waste'
    df.loc[(df['processing_unit'] == 'Pratapgarh') & ((df['agency_name'].isnull()) | (df['agency_name'] == '')), 'agency_name'] = 'SVN'
    df['net_weight'] = pd.to_numeric(df['net_weight'], errors='coerce').fillna(0)
    
    # Extract day from trip_date (DD-MM-YYYY)
    df['day'] = df['trip_date'].apply(lambda x: x.split('-')[0] if isinstance(x, str) else '')
    
    # --- 1. Overall Summary (All Units) ---
    overall_sheet = wb.create_sheet("Overall Summary")
    title = f"Municipal Corporation Faridabad - Overall {slip_type} Monthly Summary ({month_year_str})"
    overall_sheet.append([title])
    overall_sheet['A1'].font = BOLD_FONT
    overall_sheet.append([])
    
    pivot_all = pd.pivot_table(df, values='net_weight', index='day', columns='agency_name', aggfunc='sum', fill_value=0)
    
    if 'Other' in pivot_all.columns:
        pivot_all = pivot_all.drop(columns=['Other'])
        
    pivot_all['Total Weight (kg)'] = pivot_all.sum(axis=1)
    pivot_all.loc['Total'] = pivot_all.sum()
    # Add MT column (kg / 1000, rounded to 2 dp)
    pivot_all['Total Weight (MT)'] = (pivot_all['Total Weight (kg)'] / 1000).round(2)
    # Explicitly set Total row MT (pandas may not propagate new columns to .loc-set rows)
    pivot_all.loc['Total', 'Total Weight (MT)'] = round(float(pivot_all.loc['Total', 'Total Weight (kg)']) / 1000, 2)
    
    # Reorder columns by processing unit.
    # Store per_unit_agencies so we can reuse the same deduplicated list for merging.
    ordered_agencies = []
    unit_headers     = ["Day (Date)"]
    agency_headers   = [" "]
    per_unit_agencies = {}   # unit_name -> list of agencies actually used

    for unit_name, agencies in UNIT_AGENCIES.items():
        # Only add agencies not yet in ordered_agencies (avoids double-counting shared ones like Government)
        unit_agencies_in_data = [a for a in agencies
                                  if a in pivot_all.columns
                                  and a not in ("Total Weight (kg)", "Total Weight (MT)")
                                  and a not in ordered_agencies]
        per_unit_agencies[unit_name] = unit_agencies_in_data
        if unit_agencies_in_data:
            ordered_agencies.extend(unit_agencies_in_data)
            unit_headers.extend([unit_name] + [" "] * (len(unit_agencies_in_data) - 1))
            agency_headers.extend(unit_agencies_in_data)

    # Add any leftovers that might not be mapped
    leftovers = [a for a in pivot_all.columns
                 if a not in ordered_agencies
                 and a not in ("Total Weight (kg)", "Total Weight (MT)")]
    if leftovers:
        ordered_agencies.extend(leftovers)
        unit_headers.extend(["Other"] + [" "] * (len(leftovers) - 1))
        agency_headers.extend(leftovers)

    ordered_agencies.append("Total Weight (kg)")
    unit_headers.append("Total Weight (kg)")
    agency_headers.append(" ")
    ordered_agencies.append("Total Weight (MT)")
    unit_headers.append("Total Weight (MT)")
    agency_headers.append(" ")

    # Reorder pivot_all columns
    pivot_all = pivot_all[ordered_agencies]
    
    overall_sheet.append(unit_headers)
    overall_sheet.append(agency_headers)
    

    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    # Style FIRST to ensure every cell has the formatting applied to it properly
    for row in overall_sheet.iter_rows(min_row=3, max_row=4, min_col=1, max_col=len(unit_headers)):
        for cell in row:
            if cell.value is None or cell.value == "":
                cell.value = " "
            # Give MT column a distinct green tint so it stands out
            col_idx = cell.column
            is_mt_col = (cell.value == "Total Weight (MT)" or 
                         (cell.row == 4 and col_idx == len(unit_headers)))
            if is_mt_col:
                cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFFFF")
                cell.fill = PatternFill(start_color="FF1F7A4B", end_color="FF1F7A4B", fill_type="solid")
            else:
                cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFFFF")
                cell.fill = PatternFill(start_color="FF366092", end_color="FF366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(
                left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin")
            )

    from openpyxl.utils import get_column_letter
    from openpyxl.styles import PatternFill as _PF_hdr, Font as _Fnt_hdr, Alignment as _Aln_hdr, Border as _Bdr_hdr, Side as _Sd_hdr

    # 1-based column indices: row_data = [day] + ordered_agencies
    # col 1 = Day, col 2 = ordered_agencies[0], ..., col N+1 = ordered_agencies[N-1]
    # 'Total Weight (kg)'  = ordered_agencies[N-2]  → col N-2+2 = N
    # 'Total Weight (MT)'  = ordered_agencies[N-1]  → col N-1+2 = N+1
    N = len(ordered_agencies)
    kg_col_idx = N       # col of 'Total Weight (kg)'
    mt_col_idx = N + 1   # col of 'Total Weight (MT)'


    _hdr_border  = _Bdr_hdr(left=_Sd_hdr(style="thin"), right=_Sd_hdr(style="thin"),
                              top=_Sd_hdr(style="thin"),  bottom=_Sd_hdr(style="thin"))
    _center_wrap = _Aln_hdr(horizontal="center", vertical="center", wrap_text=True)

    # ── Set heading text & style BEFORE merging ──────────────────────────────
    for col_i, label, color in [
        (kg_col_idx, "Total Weight (kg)", "FF17375E"),
        (mt_col_idx, "Total Weight (MT)", "FF1F7A4B"),
    ]:
        hdr_c = overall_sheet.cell(row=3, column=col_i)
        hdr_c.value     = label
        hdr_c.font      = _Fnt_hdr(name="Arial", size=11, bold=True, color="FFFFFFFF")
        hdr_c.fill      = _PF_hdr(start_color=color, end_color=color, fill_type="solid")
        hdr_c.alignment = _center_wrap
        hdr_c.border    = _hdr_border

    # ── Merge unit-name headers (horizontal) — use per_unit_agencies to get the same count ──
    start_col = 2
    for unit_name, unit_ags in per_unit_agencies.items():
        count = len(unit_ags)
        if count > 1:
            overall_sheet.merge_cells(start_row=3, start_column=start_col,
                                       end_row=3,   end_column=start_col + count - 1)
        if count > 0:
            start_col += count

    if leftovers and len(leftovers) > 1:
        overall_sheet.merge_cells(start_row=3, start_column=start_col,
                                   end_row=3,   end_column=start_col + len(leftovers) - 1)

    # ── Merge Day, kg, MT vertically (rows 3-4) AFTER setting their heading values ──
    overall_sheet.merge_cells(start_row=3, start_column=1,          end_row=4, end_column=1)
    overall_sheet.merge_cells(start_row=3, start_column=kg_col_idx, end_row=4, end_column=kg_col_idx)
    overall_sheet.merge_cells(start_row=3, start_column=mt_col_idx, end_row=4, end_column=mt_col_idx)

    # ── Data rows ──────────────────────────────────────────────────────────────
    from openpyxl.styles import PatternFill as _PF
    MT_FILL      = _PF(start_color="FFE2EFDA", end_color="FFE2EFDA", fill_type="solid")
    MT_BOLD_FILL = _PF(start_color="FF1F7A4B", end_color="FF1F7A4B", fill_type="solid")

    current_row = 5
    for day in pivot_all.index:
        row_data = [day] + [pivot_all.loc[day, a] for a in ordered_agencies]
        overall_sheet.append(row_data)
        apply_data_row_formatting(overall_sheet, current_row)

        if day == 'Total':
            for col in range(1, mt_col_idx + 1):
                c = overall_sheet.cell(row=current_row, column=col)
                if col == mt_col_idx:
                    c.font = _Fnt_hdr(name="Arial", size=11, bold=True, color="FFFFFFFF")
                    c.fill = MT_BOLD_FILL
                else:
                    c.font = BOLD_FONT
        else:
            overall_sheet.cell(row=current_row, column=mt_col_idx).fill = MT_FILL
        current_row += 1

    # ── Explicitly write MT Total (most reliable — avoids pandas NaN propagation issues) ──
    total_row_num = current_row - 1   # current_row was incremented after Total row was written
    # Read kg total from the sheet cell that was already written (guaranteed numeric)
    kg_total_cell = overall_sheet.cell(row=total_row_num, column=kg_col_idx)
    kg_total_val  = kg_total_cell.value or 0
    mt_total_val  = round(float(kg_total_val) / 1000, 2)
    from openpyxl.styles import Alignment as _Aln_t, Border as _Bdr_t, Side as _Sd_t
    mt_total_cell = overall_sheet.cell(row=total_row_num, column=mt_col_idx)
    mt_total_cell.value     = mt_total_val
    mt_total_cell.font      = _Fnt_hdr(name="Arial", size=11, bold=True, color="FFFFFFFF")
    mt_total_cell.fill      = MT_BOLD_FILL
    mt_total_cell.alignment = _Aln_t(horizontal="center", vertical="center")
    mt_total_cell.border    = _Bdr_t(left=_Sd_t(style="thin"), right=_Sd_t(style="thin"),
                                      top=_Sd_t(style="thin"),  bottom=_Sd_t(style="thin"))

    widths = [15] + [20] * len(ordered_agencies)
    for i, width in enumerate(widths):
        overall_sheet.column_dimensions[get_column_letter(i + 1)].width = width


    # --- 2. Unit-Specific Summaries ---
    units = df['processing_unit'].unique()
    
    for unit in ['Pratapgarh', 'Mujeri', 'Gurgaon Paper Mills']:
        if unit not in units:
            continue
            
        unit_df = df[df['processing_unit'] == unit].copy()
        sheet = wb.create_sheet(f"{unit[:15]} Summary")
        
        title = f"Municipal Corporation Faridabad - {unit} - {slip_type} Monthly Summary ({month_year_str})"
        sheet.append([title])
        sheet['A1'].font = BOLD_FONT
        sheet.append([])
        
        # Pivot table: Rows = day, Columns = agency, Values = sum(net_weight)
        pivot = pd.pivot_table(unit_df, values='net_weight', index='day', columns='agency_name', aggfunc='sum', fill_value=0)
        
        if 'Other' in pivot.columns:
            pivot = pivot.drop(columns=['Other'])
        
        # Add Total kg column
        pivot['Total Weight (kg)'] = pivot.sum(axis=1)
        
        # Add Total row
        pivot.loc['Total'] = pivot.sum()
        
        # Add MT column (kg / 1000, rounded to 2 dp)
        pivot['Total Weight (MT)'] = (pivot['Total Weight (kg)'] / 1000).round(2)
        # Explicitly set Total row MT (pandas may not propagate new columns to .loc-set rows)
        pivot.loc['Total', 'Total Weight (MT)'] = round(float(pivot.loc['Total', 'Total Weight (kg)']) / 1000, 2)
        
        agencies = list(pivot.columns)
        headers = ["Day (Date)"] + agencies
        sheet.append(headers)
        
        if unit == "Pratapgarh":
            header_fill = HEADER_FILL_PRATAPGARH
        elif unit == "Mujeri":
            header_fill = HEADER_FILL_MUJERI
        else:
            from excel.formatting import HEADER_FILL_GURGAON
            header_fill = HEADER_FILL_GURGAON
            
        apply_header_formatting(sheet, 3, header_fill)

        from openpyxl.styles import PatternFill as _PF2, Font as _Font2, Alignment as _Aln2, Border as _Bdr2, Side as _Sd2
        _MT_FILL      = _PF2(start_color="FFE2EFDA", end_color="FFE2EFDA", fill_type="solid")
        _MT_BOLD_FILL = _PF2(start_color="FF1F7A4B", end_color="FF1F7A4B", fill_type="solid")
        _mt_col_idx   = len(headers)  # 1-based index of MT column
        _kg_col_idx   = len(headers) - 1  # 1-based index of kg column
        _u_center     = _Aln2(horizontal="center", vertical="center", wrap_text=True)
        _u_border     = _Bdr2(left=_Sd2(style="thin"), right=_Sd2(style="thin"),
                               top=_Sd2(style="thin"), bottom=_Sd2(style="thin"))

        # Explicitly style kg header — dark navy so it stands out as a totals column
        _kg_hdr = sheet.cell(row=3, column=_kg_col_idx)
        _kg_hdr.value     = "Total Weight (kg)"
        _kg_hdr.font      = _Font2(name="Arial", size=11, bold=True, color="FFFFFFFF")
        _kg_hdr.fill      = _PF2(start_color="FF17375E", end_color="FF17375E", fill_type="solid")
        _kg_hdr.alignment = _u_center
        _kg_hdr.border    = _u_border

        # Explicitly style MT header — green
        _mt_hdr = sheet.cell(row=3, column=_mt_col_idx)
        _mt_hdr.value     = "Total Weight (MT)"
        _mt_hdr.font      = _Font2(name="Arial", size=11, bold=True, color="FFFFFFFF")
        _mt_hdr.fill      = _MT_BOLD_FILL
        _mt_hdr.alignment = _u_center
        _mt_hdr.border    = _u_border

        current_row = 4
        for day in pivot.index:
            row_data = [day]
            for agency in agencies:
                row_data.append(pivot.loc[day, agency])
            sheet.append(row_data)
            apply_data_row_formatting(sheet, current_row)
            
            # Make Total row bold; give MT cell a green fill
            if day == 'Total':
                for col in range(1, len(headers) + 1):
                    cell = sheet.cell(row=current_row, column=col)
                    if col == _mt_col_idx:
                        cell.font = _Font2(name="Arial", size=11, bold=True, color="FFFFFFFF")
                        cell.fill = _MT_BOLD_FILL
                    else:
                        cell.font = BOLD_FONT
            else:
                sheet.cell(row=current_row, column=_mt_col_idx).fill = _MT_FILL
                    
            current_row += 1
        # Explicitly write MT Total for the unit sheet (avoids pandas NaN propagation)
        total_row_num = current_row - 1
        kg_total_cell = sheet.cell(row=total_row_num, column=_kg_col_idx)
        kg_total_val  = kg_total_cell.value or 0
        mt_total_val  = round(float(kg_total_val) / 1000, 2)
        
        mt_total_cell = sheet.cell(row=total_row_num, column=_mt_col_idx)
        mt_total_cell.value     = mt_total_val
        mt_total_cell.font      = _Font2(name="Arial", size=11, bold=True, color="FFFFFFFF")
        mt_total_cell.fill      = _MT_BOLD_FILL
        mt_total_cell.alignment = _u_center
        mt_total_cell.border    = _u_border
            
        # Set column widths
        from openpyxl.utils import get_column_letter as _gcl
        widths = [15] + [20] * len(agencies)
        for i, width in enumerate(widths):
            sheet.column_dimensions[_gcl(i + 1)].width = width

    if len(wb.sheetnames) > 1 and default_sheet.title in wb.sheetnames:
        wb.remove(default_sheet)
        
    wb.save(report_path)
    return str(report_path)
